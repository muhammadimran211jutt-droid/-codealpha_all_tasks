import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.drawing.image import Image
import os

# ==========================================
# FILE PATHS
# ==========================================

input_file = r"C:\Users\imran\Desktop\amazon-scrapper\spoon\Amazon_spoon_Cleaned.xlsx"

output_file = r"C:\Users\imran\Desktop\amazon-scrapper\spoon\Amazon_spoon_Visualization.xlsx"

chart_folder = r"C:\Users\imran\Desktop\amazon-scrapper\spoon\Charts"

os.makedirs(chart_folder, exist_ok=True)


# ==========================================
# LOAD DATA
# ==========================================

df = pd.read_excel(input_file)

print("Dataset loaded successfully!")
print("Rows:", len(df))
print("Columns:", len(df.columns))


# ==========================================
# CLEAN NUMERIC COLUMNS
# ==========================================

if "Price" in df.columns:
    df["Price"] = pd.to_numeric(df["Price"], errors="coerce")

if "Rating" in df.columns:
    df["Rating"] = pd.to_numeric(df["Rating"], errors="coerce")

if "Review Count" in df.columns:
    df["Review Count"] = pd.to_numeric(
        df["Review Count"],
        errors="coerce"
    )


# ==========================================
# CREATE EXCEL WITH DATASET
# ==========================================

with pd.ExcelWriter(
    output_file,
    engine="openpyxl"
) as writer:

    df.to_excel(
        writer,
        sheet_name="Cleaned_Data",
        index=False
    )


# ==========================================
# OPEN EXCEL
# ==========================================

wb = load_workbook(output_file)


# ==========================================
# 1. TOP 10 BRANDS
# ==========================================

if "Brand" in df.columns:

    top_brands = (
        df["Brand"]
        .dropna()
        .astype(str)
        .value_counts()
        .head(10)
    )

    chart_path = os.path.join(
        chart_folder,
        "Top_10_Brands.png"
    )

    plt.figure(figsize=(10, 6))
    top_brands.plot(kind="bar")

    plt.title("Top 10 Brands by Number of Products")
    plt.xlabel("Brand")
    plt.ylabel("Number of Products")
    plt.xticks(rotation=45, ha="right")
    plt.tight_layout()

    plt.savefig(chart_path, dpi=300)
    plt.close()

    ws = wb.create_sheet("Top_Brands")

    ws["A1"] = "Brand"
    ws["B1"] = "Product Count"

    for i, (brand, count) in enumerate(
        top_brands.items(),
        start=2
    ):
        ws.cell(i, 1, brand)
        ws.cell(i, 2, count)

    img = Image(chart_path)
    img.width = 700
    img.height = 400

    ws.add_image(img, "D2")


# ==========================================
# 2. PRICE DISTRIBUTION
# ==========================================

if "Price" in df.columns:

    price_data = df["Price"].dropna()

    chart_path = os.path.join(
        chart_folder,
        "Price_Distribution.png"
    )

    plt.figure(figsize=(10, 6))

    plt.hist(
        price_data,
        bins=20
    )

    plt.title("Price Distribution")
    plt.xlabel("Price (USD)")
    plt.ylabel("Number of Products")
    plt.tight_layout()

    plt.savefig(chart_path, dpi=300)
    plt.close()

    ws = wb.create_sheet("Price_Distribution")

    ws["A1"] = "Price (USD)"

    for i, price in enumerate(
        price_data,
        start=2
    ):
        ws.cell(i, 1, price)

    img = Image(chart_path)
    img.width = 700
    img.height = 400

    ws.add_image(img, "D2")


# ==========================================
# 3. RATING DISTRIBUTION
# ==========================================

if "Rating" in df.columns:

    rating_counts = (
        df["Rating"]
        .dropna()
        .value_counts()
        .sort_index()
    )

    chart_path = os.path.join(
        chart_folder,
        "Rating_Distribution.png"
    )

    plt.figure(figsize=(10, 6))

    rating_counts.plot(kind="bar")

    plt.title("Product Rating Distribution")
    plt.xlabel("Rating")
    plt.ylabel("Number of Products")
    plt.xticks(rotation=0)
    plt.tight_layout()

    plt.savefig(chart_path, dpi=300)
    plt.close()

    ws = wb.create_sheet("Rating_Distribution")

    ws["A1"] = "Rating"
    ws["B1"] = "Product Count"

    for i, (rating, count) in enumerate(
        rating_counts.items(),
        start=2
    ):
        ws.cell(i, 1, rating)
        ws.cell(i, 2, count)

    img = Image(chart_path)
    img.width = 700
    img.height = 400

    ws.add_image(img, "D2")


# ==========================================
# 4. PRICE VS RATING
# ==========================================

if "Price" in df.columns and "Rating" in df.columns:

    plot_data = df[
        ["Price", "Rating"]
    ].dropna()

    chart_path = os.path.join(
        chart_folder,
        "Price_vs_Rating.png"
    )

    plt.figure(figsize=(10, 6))

    plt.scatter(
        plot_data["Price"],
        plot_data["Rating"],
        alpha=0.6
    )

    plt.title("Price vs Rating")
    plt.xlabel("Price (USD)")
    plt.ylabel("Rating")
    plt.tight_layout()

    plt.savefig(chart_path, dpi=300)
    plt.close()

    ws = wb.create_sheet("Price_vs_Rating")

    ws["A1"] = "Price (USD)"
    ws["B1"] = "Rating"

    for i, row in enumerate(
        plot_data.itertuples(index=False),
        start=2
    ):
        ws.cell(i, 1, row.Price)
        ws.cell(i, 2, row.Rating)

    img = Image(chart_path)
    img.width = 700
    img.height = 400

    ws.add_image(img, "D2")


# ==========================================
# 5. TOP 10 PRODUCTS BY REVIEW COUNT
# ==========================================

if (
    "Review Count" in df.columns
    and "Product Title" in df.columns
):

    top_reviews = (
        df[
            ["Product Title", "Review Count"]
        ]
        .dropna()
        .sort_values(
            "Review Count",
            ascending=False
        )
        .head(10)
        .copy()
    )

    top_reviews["Short Title"] = (
        top_reviews["Product Title"]
        .astype(str)
        .str[:45]
    )

    chart_path = os.path.join(
        chart_folder,
        "Top_10_Products_Reviews.png"
    )

    plt.figure(figsize=(12, 7))

    plt.barh(
        top_reviews["Short Title"],
        top_reviews["Review Count"]
    )

    plt.title("Top 10 Products by Review Count")
    plt.xlabel("Number of Reviews")
    plt.ylabel("Product")

    plt.gca().invert_yaxis()

    plt.tight_layout()

    plt.savefig(chart_path, dpi=300)
    plt.close()

    ws = wb.create_sheet("Top_Reviews")

    ws["A1"] = "Product Title"
    ws["B1"] = "Review Count"

    for i, row in enumerate(
        top_reviews.iterrows(),
        start=2
    ):
        index, data = row

        ws.cell(
            i,
            1,
            data["Product Title"]
        )

        ws.cell(
            i,
            2,
            data["Review Count"]
        )

    img = Image(chart_path)
    img.width = 800
    img.height = 450

    ws.add_image(img, "D2")


# ==========================================
# SAVE FINAL EXCEL
# ==========================================

wb.save(output_file)

print("\n========================================")
print("TASK 3 COMPLETED SUCCESSFULLY!")
print("========================================")
print("Final Excel file:")
print(output_file)
print("\nYour Excel contains:")
print("1. Cleaned Dataset")
print("2. Top 10 Brands Chart")
print("3. Price Distribution Chart")
print("4. Rating Distribution Chart")
print("5. Price vs Rating Chart")
print("6. Top 10 Products by Reviews Chart")